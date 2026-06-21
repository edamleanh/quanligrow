import json
import zipfile
import re
import os
import urllib.request
import urllib.error
from http.server import BaseHTTPRequestHandler
from io import BytesIO
import openpyxl

def get_vn_sort_key(text):
    text = str(text).strip().lower()
    text = text.replace('ð', 'đ') # Handle Unicode Eth which looks like Đ
    alphabet = "a á à ả ã ạ ă ắ ằ ẳ ẵ ặ â ấ ầ ẩ ẫ ậ b c d đ e é è ẻ ẽ ẹ ê ế ề ể ễ ệ f g h i í ì ỉ ĩ ị j k l m n o ó ò ỏ õ ọ ô ố ồ ổ ỗ ộ ơ ớ ờ ở ỡ ợ p q r s t u ú ù ủ ũ ụ ư ứ ừ ử ữ ự v w x y ý ỳ ỷ ỹ ỵ z"
    order = {c: f"{i+1:03d}" for i, c in enumerate(alphabet.split())}
    result = []
    for char in text:
        if char in order:
            result.append(order[char])
        elif char.isspace():
            result.append("000")
        else:
            result.append(f"999{ord(char):05d}")
    return "-".join(result)

class handler(BaseHTTPRequestHandler):
    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type, X-Requested-With')
        self.end_headers()

    def do_POST(self):
        try:
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            students_data = json.loads(post_data)

            target_subjects = ["TOÁN", "VĂN", "Hóa", "AV", "Lý"]
            
            # Find the template in the root path or current path
            base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            template_path = os.path.join(base_dir, "file chuẩn.xlsx")
            if not os.path.exists(template_path):
                # Fallback check
                template_path = os.path.join(os.path.dirname(__file__), "file chuẩn.xlsx")

            # Fetch class status from Supabase
            class_status = {}
            try:
                url = "https://pvlyaubewwhgzhirmwgi.supabase.co/rest/v1/ds_tong?select=Ghi%20ch%C3%BA&STT=eq.99999"
                req = urllib.request.Request(url)
                req.add_header('apikey', 'sb_publishable_aBQHszEUh4jdxO7ebD9DCg_tH-Qi3kw')
                req.add_header('Authorization', 'Bearer sb_publishable_aBQHszEUh4jdxO7ebD9DCg_tH-Qi3kw')
                response = urllib.request.urlopen(req)
                res_data = json.loads(response.read().decode('utf-8'))
                if res_data and len(res_data) > 0 and res_data[0].get("Ghi chú"):
                    class_status = json.loads(res_data[0]["Ghi chú"])
            except Exception as e:
                print("Failed to fetch class status:", e)

            # Process logic
            zip_buffer = BytesIO()
            with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for subject in target_subjects:
                    students_by_class = {}
                    for row in students_data:
                        ho = row.get("HỌ", "")
                        ten = row.get("TÊN", "")
                        if not ho and not ten: continue
                        sub_val = row.get(subject, "")
                        if not sub_val: continue
                        lop = row.get("LỚP", "")
                        class_name = f"{lop}{sub_val}".strip()
                        if not class_name or len(class_name) > 4: continue
                        if class_name not in students_by_class:
                            students_by_class[class_name] = []
                        students_by_class[class_name].append({
                            "ho": ho or "",
                            "ten": ten or "",
                            "sodt": row.get("SỐ ĐT", ""),
                            "thieu": row.get("THIẾU", 0)
                        })

                    def class_sort_key(c):
                        m = re.match(r"^(1[0-2]|[1-9])(.*)$", str(c).strip())
                        if m:
                            grade = int(m.group(1))
                            suffix = m.group(2).strip().upper()
                        else:
                            return (float('inf'), 3, str(c))
                        if not suffix: priority = 0
                        elif suffix[0].isdigit(): priority = 0
                        elif suffix.startswith("O"): priority = 1
                        else: priority = 2
                        return (grade, priority, suffix)

                    active_classes = []
                    for c in students_by_class.keys():
                        full_class_name = f"{c} {subject}".strip()
                        if class_status.get(full_class_name, True) is not False:
                            active_classes.append(c)

                    classes = sorted(active_classes, key=class_sort_key)
                    blocks = []
                    for class_name in classes:
                        students = students_by_class[class_name]
                        students.sort(key=lambda x: (get_vn_sort_key(x.get("ten", "")), get_vn_sort_key(x.get("ho", ""))))
                        for chunk_idx in range(0, len(students), 25):
                            chunk = students[chunk_idx:chunk_idx+25]
                            blocks.append({
                                "class_name": class_name,
                                "students": chunk,
                                "start_idx": chunk_idx
                            })
                    
                    wb_out = openpyxl.load_workbook(template_path)
                    sheet = wb_out.active
                    
                    for i, block in enumerate(blocks):
                        class_name = block["class_name"]
                        students = block["students"]
                        stt_offset = block["start_idx"]
                        base_row = i * 30 + 1
                        
                        title_cell = sheet.cell(row=base_row, column=1)
                        if title_cell.value and isinstance(title_cell.value, str):
                            if type(title_cell).__name__ != 'MergedCell':
                                title_cell.value = f"DANH SÁCH {subject.upper()}"
                        
                        if type(sheet.cell(row=base_row, column=9)).__name__ != 'MergedCell':
                            sheet.cell(row=base_row, column=9).value = class_name
                        
                        start_row = base_row + 4
                        for idx, student in enumerate(students):
                            current_row = start_row + idx
                            sheet.cell(row=current_row, column=1).value = stt_offset + idx + 1
                            sheet.cell(row=current_row, column=2).value = student["ho"]
                            sheet.cell(row=current_row, column=3).value = student["ten"]
                            sheet.cell(row=current_row, column=4).value = student["sodt"]
                            
                        for idx in range(len(students), 25):
                            current_row = start_row + idx
                            sheet.cell(row=current_row, column=1).value = ""
                            sheet.cell(row=current_row, column=2).value = ""
                            sheet.cell(row=current_row, column=3).value = ""
                            sheet.cell(row=current_row, column=4).value = ""
                            
                    for i in range(len(blocks), 150):
                        base_row = i * 30 + 1
                        if type(sheet.cell(row=base_row, column=9)).__name__ != 'MergedCell':
                            sheet.cell(row=base_row, column=9).value = ""
                        title_cell = sheet.cell(row=base_row, column=1)
                        if type(title_cell).__name__ != 'MergedCell':
                            title_cell.value = ""

                    excel_buffer = BytesIO()
                    wb_out.save(excel_buffer)
                    excel_buffer.seek(0)
                    
                    file_to_zip = f"DanhSach_{subject.upper()}_TatCaLop.xlsx"
                    zipf.writestr(file_to_zip, excel_buffer.read())

            zip_buffer.seek(0)
            
            self.send_response(200)
            self.send_header('Access-Control-Allow-Origin', '*')
            self.send_header('Content-Type', 'application/zip')
            self.send_header('Content-Disposition', 'attachment; filename="DanhSachTruatXuat.zip"')
            self.end_headers()
            self.wfile.write(zip_buffer.read())

        except Exception as e:
            self.send_response(500)
            self.send_header('Access-Control-Allow-Origin', '*')
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(json.dumps({'error': str(e)}).encode('utf-8'))
