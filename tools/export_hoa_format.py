import openpyxl

def extract_sheet():
    file_path = 'HÈ.xlsx'
    new_file_path = 'HOA_Rieng_Giu_Dinh_Dang.xlsx'
    
    print("Loading file, this might take a moment to preserve formatting...")
    
    # Load the workbook preserving data and styles
    wb = openpyxl.load_workbook(file_path)
    
    # Find the target sheet name (case-insensitive)
    target_sheet = None
    for sheet in wb.sheetnames:
        if sheet.upper() in ['HOÁ', 'HÓA']:
            target_sheet = sheet
            break
            
    if not target_sheet:
        print("Sheet HOA not found!")
        print("Available sheets:", wb.sheetnames)
        return
        
    print(f"Found sheet: {target_sheet}. Extracting...")
    
    # Remove all other sheets
    for sheet in wb.sheetnames:
        if sheet != target_sheet:
            del wb[sheet]
            
    # Save to a new file
    print("Saving to new file...")
    wb.save(new_file_path)
    print(f"Success! Saved to {new_file_path}")

if __name__ == '__main__':
    extract_sheet()
